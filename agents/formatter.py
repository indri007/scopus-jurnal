"""
Phase 15+16 — Formatter Agent
Assembles DOCX and/or LaTeX output from manuscript sections.
"""

import io
import logging
import os
import re
import subprocess
from typing import Optional

from core.state import PaperState
import config

logger = logging.getLogger(__name__)

SECTION_ORDER = [
    "abstract",
    "introduction",
    "literature_review",
    "theoretical_framework",
    "methodology",
    "results",
    "discussion",
    "theoretical_implications",
    "practical_implications",
    "limitations",
    "conclusion",
]

SECTION_TITLES = {
    "abstract":                "Abstract",
    "introduction":            "Introduction",
    "literature_review":       "Literature Review",
    "theoretical_framework":   "Theoretical Framework",
    "methodology":             "Methodology",
    "results":                 "Results",
    "discussion":              "Discussion",
    "theoretical_implications":"Theoretical Implications",
    "practical_implications":  "Practical Implications",
    "limitations":             "Limitations",
    "conclusion":              "Conclusion",
}


def run(state: PaperState, progress_cb=None) -> PaperState:
    state.mark_running("formatting")
    fmt = state.output_format or config.OUTPUT_FORMAT
    out_dir = state.path("final")
    os.makedirs(out_dir, exist_ok=True)

    if fmt in ("docx", "both"):
        _cb(progress_cb, "Generating DOCX…")
        try:
            docx_path = _generate_docx(state, out_dir)
            _cb(progress_cb, f"  ✅ DOCX: {docx_path}")
        except Exception as e:
            logger.error("DOCX generation failed: %s", e)
            state.warnings.append(f"DOCX generation failed: {e}")

    if fmt in ("latex", "both"):
        _cb(progress_cb, "Generating LaTeX…")
        try:
            tex_path = _generate_latex(state, out_dir)
            _cb(progress_cb, f"  ✅ LaTeX: {tex_path}")

            _cb(progress_cb, "Compiling PDF…")
            pdf_path = _compile_pdf(tex_path, out_dir)
            if pdf_path:
                _cb(progress_cb, f"  ✅ PDF: {pdf_path}")
            else:
                _cb(progress_cb, "  ⚠️ PDF compilation failed — LaTeX source available")
        except Exception as e:
            logger.error("LaTeX generation failed: %s", e)
            state.warnings.append(f"LaTeX generation failed: {e}")

    state.mark_done("formatting", out_dir)
    state.save()
    _cb(progress_cb, "✅ Formatting complete")
    return state


# ── DOCX ──────────────────────────────────────────────────────────────────────

def _generate_docx(state: PaperState, out_dir: str) -> str:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)

    # Title
    title_para = doc.add_heading(state.topic, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Authors / journal line
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if state.target_journal:
        meta.add_run(f"Prepared for: {state.target_journal}").italic = True

    doc.add_paragraph()  # spacer

    # Sections
    for sec_key in SECTION_ORDER:
        text = state.manuscript.get(sec_key, "")
        if not text:
            continue
        title = SECTION_TITLES.get(sec_key, sec_key.replace("_", " ").title())
        doc.add_heading(title, level=1)
        _add_markdown_text(doc, text)

    # References
    doc.add_heading("References", level=1)
    bib_path = state.path("manuscript/references.bib")
    if os.path.exists(bib_path):
        with open(bib_path, "r", encoding="utf-8") as f:
            bib_text = f.read()
        _add_bib_as_apa(doc, bib_text)
    else:
        doc.add_paragraph("[EVIDENCE_REQUIRED: references.bib not found]")

    out_path = os.path.join(out_dir, "manuscript_final.docx")
    doc.save(out_path)
    return out_path


def _add_markdown_text(doc, text: str):
    """Convert basic markdown to DOCX paragraphs."""
    from docx.shared import Pt
    for line in text.split("\n"):
        line = line.rstrip()
        if not line:
            doc.add_paragraph()
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:])
        else:
            p = doc.add_paragraph()
            # Handle **bold** and *italic*
            _add_formatted_run(p, line)


def _add_formatted_run(para, text: str):
    """Parse **bold** and *italic* markers into runs."""
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = para.add_run(part[1:-1])
            run.italic = True
        else:
            para.add_run(part)


def _add_bib_as_apa(doc, bib_text: str):
    """Convert BibTeX entries to APA-style paragraphs."""
    entries = re.split(r'\n@', bib_text)
    for entry in entries:
        entry = entry.strip()
        if not entry:
            continue
        if not entry.startswith("@"):
            entry = "@" + entry
        author = _bib_field(entry, "author") or "Unknown"
        year   = _bib_field(entry, "year") or "n.d."
        title  = _bib_field(entry, "title") or ""
        journal= _bib_field(entry, "journal") or ""
        doi    = _bib_field(entry, "doi") or ""
        apa = f"{author} ({year}). {title}. {journal}."
        if doi:
            apa += f" https://doi.org/{doi}"
        p = doc.add_paragraph(apa, style="Normal")
        p.paragraph_format.left_indent = __import__("docx.shared", fromlist=["Inches"]).Inches(0.5)


def _bib_field(entry: str, field: str) -> Optional[str]:
    m = re.search(rf'{field}\s*=\s*\{{(.*?)\}}', entry, re.DOTALL | re.IGNORECASE)
    return m.group(1).strip() if m else None


# ── LaTeX ─────────────────────────────────────────────────────────────────────

def _generate_latex(state: PaperState, out_dir: str) -> str:
    bib_path = state.path("manuscript/references.bib")
    bib_exists = os.path.exists(bib_path)

    # Copy bib to final dir
    if bib_exists:
        import shutil
        shutil.copy(bib_path, os.path.join(out_dir, "references.bib"))

    # Copy figures
    figs_src = state.path("figures")
    if os.path.exists(figs_src):
        import shutil
        figs_dst = os.path.join(out_dir, "figures")
        if not os.path.exists(figs_dst):
            shutil.copytree(figs_src, figs_dst)

    sections_latex = []
    for sec_key in SECTION_ORDER:
        text = state.manuscript.get(sec_key, "")
        if not text:
            continue
        title = SECTION_TITLES.get(sec_key, sec_key.replace("_", " ").title())
        if sec_key == "abstract":
            sections_latex.append(f"\\begin{{abstract}}\n{_md_to_latex(text)}\n\\end{{abstract}}")
        else:
            sections_latex.append(f"\\section{{{title}}}\n{_md_to_latex(text)}")

    # References section
    if bib_exists:
        ref_block = "\\bibliographystyle{apalike}\n\\bibliography{references}"
    else:
        ref_block = "% [EVIDENCE_REQUIRED: references.bib not generated]"

    citation_style_pkg = {
        "apa":  "\\usepackage[style=apa,backend=biber]{biblatex}",
        "ieee": "\\usepackage[style=ieee,backend=biber]{biblatex}",
        "acm":  "\\usepackage[style=acm,backend=biber]{biblatex}",
    }.get(state.citation_style, "\\usepackage[style=apa,backend=biber]{biblatex}")

    tex = rf"""\documentclass[12pt,a4paper]{{article}}
\usepackage[utf8]{{inputenc}}
\usepackage[T1]{{fontenc}}
\usepackage{{lmodern}}
\usepackage{{microtype}}
\usepackage{{geometry}}
\geometry{{margin=2.5cm}}
\usepackage{{hyperref}}
\hypersetup{{colorlinks=true,linkcolor=blue,citecolor=blue,urlcolor=blue}}
\usepackage{{graphicx}}
\usepackage{{booktabs}}
\usepackage{{amsmath}}
{citation_style_pkg}
\usepackage{{setspace}}
\doublespacing

\title{{{_escape_latex(state.topic)}}}
\author{{}}
\date{{}}

\begin{{document}}

\maketitle

{chr(10).join(sections_latex)}

\section*{{References}}
{ref_block}

\end{{document}}
"""
    tex_path = os.path.join(out_dir, "manuscript_final.tex")
    with open(tex_path, "w", encoding="utf-8") as f:
        f.write(tex)
    return tex_path


def _compile_pdf(tex_path: str, out_dir: str) -> Optional[str]:
    """Try to compile LaTeX to PDF. Returns PDF path or None."""
    for engine in ("pdflatex", "xelatex"):
        try:
            result = subprocess.run(
                [engine, "-interaction=nonstopmode", "-output-directory", out_dir, tex_path],
                capture_output=True, text=True, timeout=60
            )
            pdf_path = tex_path.replace(".tex", ".pdf")
            if os.path.exists(pdf_path):
                # Run twice for cross-references
                subprocess.run(
                    [engine, "-interaction=nonstopmode", "-output-directory", out_dir, tex_path],
                    capture_output=True, text=True, timeout=60
                )
                return pdf_path
        except (FileNotFoundError, subprocess.TimeoutExpired):
            continue
    return None


def _md_to_latex(text: str) -> str:
    """Convert basic markdown to LaTeX."""
    lines = text.split("\n")
    out = []
    for line in lines:
        # Headings
        if line.startswith("### "):
            out.append(f"\\subsubsection{{{_escape_latex(line[4:])}}}")
        elif line.startswith("## "):
            out.append(f"\\subsection{{{_escape_latex(line[3:])}}}")
        elif line.startswith("# "):
            out.append(f"\\section{{{_escape_latex(line[2:])}}}")
        # Bullet lists
        elif line.startswith("- ") or line.startswith("* "):
            out.append(f"\\item {_escape_latex(line[2:])}")
        else:
            line = re.sub(r'\*\*(.*?)\*\*', r'\\textbf{\1}', line)
            line = re.sub(r'\*(.*?)\*',   r'\\textit{\1}', line)
            line = _escape_latex(line, skip_commands=True)
            out.append(line)
    # Wrap consecutive \item lines in itemize
    result = []
    in_list = False
    for line in out:
        if line.startswith("\\item "):
            if not in_list:
                result.append("\\begin{itemize}")
                in_list = True
            result.append(line)
        else:
            if in_list:
                result.append("\\end{itemize}")
                in_list = False
            result.append(line)
    if in_list:
        result.append("\\end{itemize}")
    return "\n".join(result)


def _escape_latex(text: str, skip_commands: bool = False) -> str:
    if skip_commands:
        return text
    replacements = [
        ("&", "\\&"), ("%", "\\%"), ("$", "\\$"), ("#", "\\#"),
        ("_", "\\_"), ("{", "\\{"), ("}", "\\}"), ("~", "\\textasciitilde{}"),
        ("^", "\\textasciicircum{}"),
    ]
    for old, new in replacements:
        text = text.replace(old, new)
    return text


def _cb(fn, msg: str):
    if fn:
        fn(msg)
    else:
        logger.info(msg)
